"""Extrai a tabela de comprimentos equivalentes da planilha Excel para JSON.

Uso (a partir da raiz do projeto):
    python scripts/extrair_planilha.py
    python scripts/extrair_planilha.py --planilha caminho.xlsx --saida dados/comprimentos_equivalentes.json

A planilha guarda os valores como texto com vírgula decimal ("0,35"). Cada
valor é gravado no JSON como texto decimal com ponto ("0.35"), para que o
backend o leia como Decimal exato. Combinações inexistentes simplesmente não
são gravadas: nenhuma lacuna vira zero.
"""

import argparse
import json
import re
import sys
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
PLANILHA_PADRAO = RAIZ / "dados" / "fonte" / "1 Perda de Carga Localizada.xlsx"
SAIDA_PADRAO = RAIZ / "dados" / "comprimentos_equivalentes.json"
IMAGENS_PADRAO = RAIZ / "static" / "img" / "componentes"
PREFIXO_URL_IMAGENS = "img/componentes"

ABA = "Plan1"
CABECALHOS_FIXOS = {"A": "Descrição", "B": "Foto", "C": "Material"}
COLUNA_DESCRICAO, COLUNA_MATERIAL, PRIMEIRA_COLUNA_DIAMETRO = 1, 3, 4

UNIDADE = "m"
UNIDADE_ORIGEM = (
    "Tabela de referência reproduzida nos documentos da disciplina: "
    "'Diâmetro nominal X Equivalência em metros de canalização'."
)

# Diâmetros nominais aceitos: rótulo da planilha -> (id, polegadas, DN em mm).
# O DN é a designação nominal (ISO 6708 / ABNT); não é diâmetro interno nem externo.
DIAMETROS_CONHECIDOS = {
    '1/4"': ("1/4", "0.25", 8),
    '3/8"': ("3/8", "0.375", 10),
    '1/2"': ("1/2", "0.5", 15),
    '3/4"': ("3/4", "0.75", 20),
    '1"': ("1", "1", 25),
    '1 1/4"': ("1-1/4", "1.25", 32),
    '1 1/2"': ("1-1/2", "1.5", 40),
    '2"': ("2", "2", 50),
    '2 1/2"': ("2-1/2", "2.5", 65),
    '3"': ("3", "3", 80),
    '4"': ("4", "4", 100),
    '5"': ("5", "5", 125),
    '6"': ("6", "6", 150),
    '8"': ("8", "8", 200),
    '10"': ("10", "10", 250),
    '12"': ("12", "12", 300),
    '14"': ("14", "14", 350),
}

# Componentes cuja quantidade é um comprimento de tubo em metros (aceita
# decimais). Para os demais, a quantidade é um número inteiro de peças.
COMPONENTES_MEDIDOS_EM_METROS = {"trecho-reto"}


class ErroPlanilha(Exception):
    pass


def normalizar_texto(texto):
    """Junta quebras de linha e espaços repetidos sem alterar as palavras."""
    return re.sub(r"\s+", " ", str(texto)).strip()


def gerar_id(texto):
    sem_acento = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", sem_acento.lower()).strip("-")


def converter_valor(bruto, celula):
    """Converte '0,35' ou 1 em texto decimal exato ('0.35', '1')."""
    if bruto is None or (isinstance(bruto, str) and not bruto.strip()):
        return None
    if isinstance(bruto, bool):
        raise ErroPlanilha(f"{celula}: valor booleano inesperado")
    if isinstance(bruto, float):
        # Não ocorre na planilha atual; repr evita dígitos espúrios de ponto flutuante.
        texto = repr(bruto)
    else:
        texto = str(bruto).strip().replace(",", ".")
    try:
        valor = Decimal(texto)
    except InvalidOperation:
        raise ErroPlanilha(f"{celula}: valor não numérico {bruto!r}") from None
    if not valor.is_finite() or valor <= 0:
        raise ErroPlanilha(f"{celula}: valor fora do esperado {bruto!r}")
    return str(valor)


def ler_diametros(aba):
    diametros = []
    for coluna in range(PRIMEIRA_COLUNA_DIAMETRO, aba.max_column + 1):
        celula = aba.cell(row=1, column=coluna)
        rotulo = normalizar_texto(celula.value or "")
        if rotulo not in DIAMETROS_CONHECIDOS:
            raise ErroPlanilha(f"{celula.coordinate}: diâmetro desconhecido {celula.value!r}")
        id_, polegadas, dn = DIAMETROS_CONHECIDOS[rotulo]
        diametros.append({
            "id": id_,
            "rotulo": rotulo,
            "polegadas": polegadas,
            "dn_mm": dn,
            "coluna": celula.column_letter,
        })
    return diametros


def conferir_cabecalho(aba):
    for letra, esperado in CABECALHOS_FIXOS.items():
        encontrado = normalizar_texto(aba[f"{letra}1"].value or "")
        if encontrado != esperado:
            raise ErroPlanilha(f"{letra}1: esperado {esperado!r}, encontrado {encontrado!r}")


_NS = {
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}
_ATRIBUTO_RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
_ATRIBUTO_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def _relacoes(pacote, caminho_parte):
    """Lê o arquivo .rels de uma parte do pacote: {rId: caminho absoluto no zip}."""
    pasta, nome = caminho_parte.rsplit("/", 1)
    caminho_rels = f"{pasta}/_rels/{nome}.rels"
    if caminho_rels not in pacote.namelist():
        return {}
    relacoes = {}
    for rel in ET.fromstring(pacote.read(caminho_rels)).findall("pr:Relationship", _NS):
        alvo = rel.get("Target")
        partes = (alvo.lstrip("/") if alvo.startswith("/") else f"{pasta}/{alvo}").split("/")
        normalizado = []
        for parte in partes:
            if parte == "..":
                normalizado.pop()
            elif parte not in ("", "."):
                normalizado.append(parte)
        relacoes[rel.get("Id")] = "/".join(normalizado)
    return relacoes


def ler_imagens_por_linha(caminho_planilha):
    """Fotos da coluna B da aba, indexadas pela linha da planilha: {linha: (bytes, extensão)}.

    Lido direto do XML do pacote .xlsx, porque o openpyxl só carrega imagens
    quando o Pillow está instalado.
    """
    with zipfile.ZipFile(caminho_planilha) as pacote:
        livro = ET.fromstring(pacote.read("xl/workbook.xml"))
        abas = {a.get("name"): a.get(_ATRIBUTO_RID) for a in livro.find("m:sheets", _NS)}
        caminho_aba = _relacoes(pacote, "xl/workbook.xml")[abas[ABA]]
        aba = ET.fromstring(pacote.read(caminho_aba))
        desenho = aba.find("m:drawing", _NS)
        if desenho is None:
            return {}
        caminho_desenho = _relacoes(pacote, caminho_aba)[desenho.get(_ATRIBUTO_RID)]
        midias = _relacoes(pacote, caminho_desenho)

        imagens = {}
        for ancora in ET.fromstring(pacote.read(caminho_desenho)):
            inicio = ancora.find("xdr:from", _NS)
            blip = ancora.find(".//xdr:pic//a:blip", _NS)
            if inicio is None or blip is None:
                continue  # formas sem imagem
            linha = int(inicio.find("xdr:row", _NS).text) + 1
            if linha in imagens:
                raise ErroPlanilha(f"linha {linha}: mais de uma imagem")
            caminho_midia = midias[blip.get(_ATRIBUTO_EMBED)]
            imagens[linha] = (pacote.read(caminho_midia), Path(caminho_midia).suffix.lower())
        return imagens


def extrair(caminho_planilha, pasta_imagens=None):
    """Monta a base de referência. Com pasta_imagens, grava também as fotos dos componentes."""
    livro = openpyxl.load_workbook(caminho_planilha, read_only=False, data_only=True)
    if ABA not in livro.sheetnames:
        raise ErroPlanilha(f"aba {ABA!r} não encontrada (abas: {livro.sheetnames})")
    aba = livro[ABA]
    conferir_cabecalho(aba)
    diametros = ler_diametros(aba)
    imagens = ler_imagens_por_linha(caminho_planilha)
    if pasta_imagens is not None:
        Path(pasta_imagens).mkdir(parents=True, exist_ok=True)
        for antiga in Path(pasta_imagens).iterdir():
            antiga.unlink()

    componentes = {}
    materiais = {}
    lacunas = []
    for linha in range(2, aba.max_row + 1):
        descricao_bruta = aba.cell(row=linha, column=COLUNA_DESCRICAO).value
        material_bruto = aba.cell(row=linha, column=COLUNA_MATERIAL).value
        if descricao_bruta is None and material_bruto is None:
            continue
        if descricao_bruta is None or material_bruto is None:
            raise ErroPlanilha(f"linha {linha}: descrição ou material ausente")

        nome = normalizar_texto(descricao_bruta)
        material_rotulo = normalizar_texto(material_bruto)
        material_id = gerar_id(material_rotulo)
        materiais.setdefault(material_id, {"id": material_id, "rotulo": material_rotulo})

        componente_id = gerar_id(nome)
        componente = componentes.setdefault(componente_id, {
            "id": componente_id,
            "nome": nome,
            "unidade_quantidade": "m" if componente_id in COMPONENTES_MEDIDOS_EM_METROS else "un",
            "valores": {},
            "celulas": {},
            "imagens": {},
        })
        if componente["nome"] != nome:
            raise ErroPlanilha(f"linha {linha}: id {componente_id!r} colide com {componente['nome']!r}")
        if material_id in componente["valores"]:
            raise ErroPlanilha(f"linha {linha}: {nome} / {material_rotulo} duplicado")

        valores, celulas = {}, {}
        for diametro in diametros:
            celula = f"{diametro['coluna']}{linha}"
            valor = converter_valor(aba[celula].value, celula)
            if valor is None:
                lacunas.append(celula)
                continue
            valores[diametro["id"]] = valor
            celulas[diametro["id"]] = celula
        componente["valores"][material_id] = valores
        componente["celulas"][material_id] = celulas

        if linha in imagens:
            conteudo, extensao = imagens[linha]
            nome_arquivo = f"{componente_id}-{material_id}{extensao}"
            componente["imagens"][material_id] = f"{PREFIXO_URL_IMAGENS}/{nome_arquivo}"
            if pasta_imagens is not None:
                (Path(pasta_imagens) / nome_arquivo).write_bytes(conteudo)

    return {
        "metadados": {
            "fonte": Path(caminho_planilha).name,
            "aba": ABA,
            "unidade": UNIDADE,
            "origem_da_unidade": UNIDADE_ORIGEM,
            "representacao_dos_valores": (
                "Texto decimal exato com ponto (ex.: '0.35'), lido como Decimal pelo backend. "
                "Combinações ausentes na planilha não aparecem em 'valores'."
            ),
            "celulas_vazias_na_fonte": lacunas,
            "extraido_em": date.today().isoformat(),
        },
        "materiais": list(materiais.values()),
        "diametros": [{k: v for k, v in d.items() if k != "coluna"} for d in diametros],
        "componentes": list(componentes.values()),
    }


def main(argumentos=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--planilha", type=Path, default=PLANILHA_PADRAO)
    parser.add_argument("--saida", type=Path, default=SAIDA_PADRAO)
    parser.add_argument("--imagens", type=Path, default=IMAGENS_PADRAO)
    args = parser.parse_args(argumentos)
    try:
        base = extrair(args.planilha, args.imagens)
    except (ErroPlanilha, FileNotFoundError, KeyError, zipfile.BadZipFile) as erro:
        print(f"Erro ao extrair a planilha: {erro}", file=sys.stderr)
        return 1
    args.saida.parent.mkdir(parents=True, exist_ok=True)
    args.saida.write_text(json.dumps(base, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    total = sum(len(v) for c in base["componentes"] for v in c["valores"].values())
    print(f"{len(base['componentes'])} componentes, {len(base['materiais'])} materiais, "
          f"{len(base['diametros'])} diâmetros, {total} valores -> {args.saida.name}")
    fotos = sum(len(c["imagens"]) for c in base["componentes"])
    print(f"{fotos} imagens -> {args.imagens.name}/")
    return 0


if __name__ == "__main__":
    sys.exit(main())
