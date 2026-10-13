"""Confere o JSON contra todas as células da planilha original."""

import json
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from scripts.extrair_planilha import PLANILHA_PADRAO, SAIDA_PADRAO, extrair

pytestmark = pytest.mark.skipif(not Path(PLANILHA_PADRAO).exists(), reason="planilha de origem ausente")


def test_json_versionado_corresponde_a_extracao_atual():
    gerado = extrair(PLANILHA_PADRAO)
    salvo = json.loads(Path(SAIDA_PADRAO).read_text(encoding="utf-8"))
    for chave in ("materiais", "diametros", "componentes"):
        assert gerado[chave] == salvo[chave]


def test_cada_valor_do_json_bate_com_a_celula_de_origem():
    aba = openpyxl.load_workbook(PLANILHA_PADRAO)["Plan1"]
    base = json.loads(Path(SAIDA_PADRAO).read_text(encoding="utf-8"))
    conferidos = 0
    for componente in base["componentes"]:
        for material, valores in componente["valores"].items():
            for diametro, valor in valores.items():
                celula = componente["celulas"][material][diametro]
                original = aba[celula].value
                assert Decimal(str(original).replace(",", ".")) == Decimal(valor), celula
                linha = aba[celula].row
                assert " ".join(str(aba.cell(row=linha, column=1).value).split()) == componente["nome"]
                assert str(aba.cell(row=linha, column=3).value).strip().lower() == material
                conferidos += 1
    assert conferidos == 64 * 17


def test_amostra_manual_de_celulas():
    base = json.loads(Path(SAIDA_PADRAO).read_text(encoding="utf-8"))
    por_id = {c["id"]: c for c in base["componentes"]}
    assert por_id["joelho-90"]["valores"]["pvc"]["3/4"] == "1.20"          # G2
    assert por_id["joelho-90"]["valores"]["metal"]["1"] == "0.94"          # H3
    assert por_id["curva-de-transposicao"]["valores"]["metal"]["1/2"] == "0.87"  # F21, preservado
    assert por_id["registro-de-globo"]["valores"]["metal"]["14"] == "120.90"     # T58
    assert por_id["trecho-reto"]["valores"]["pvc"]["2"] == "1"             # K64
    assert "pvc" not in por_id["registro-de-globo"]["valores"]
    assert "metal" not in por_id["filtro-y"]["valores"]
